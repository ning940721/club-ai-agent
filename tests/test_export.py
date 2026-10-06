from datetime import date, time

import pytest

from club_agent.agenda import AgendaRow, schedule, schedule_markdown, template_rows, total_minutes
from club_agent.export import inline_runs, parse_markdown, pdf_available, to_docx, to_pdf

SAMPLE = """# 測試社 進度彙整

> **問題：** 預算怎麼分？

## 行動步驟
| # | 步驟 | 負責 |
|---|---|---|
| 1 | 先做**預算表** | 財務 |

- 第一點
1. 編號一
一般段落，含 **粗體** 與 *斜體*。
"""


def test_parse_markdown_blocks():
    kinds = [b.kind for b in parse_markdown(SAMPLE)]
    assert kinds == ["heading", "quote", "heading", "table", "bullet", "number", "para"]
    table = next(b for b in parse_markdown(SAMPLE) if b.kind == "table")
    assert table.rows == [["#", "步驟", "負責"], ["1", "先做**預算表**", "財務"]]


def test_inline_runs_bold_and_strips_italic():
    assert inline_runs("含 **粗體** 與 *斜體*") == [("含 ", False), ("粗體", True), (" 與 斜體", False)]


def test_docx_contains_text():
    from io import BytesIO

    from docx import Document

    doc = Document(BytesIO(to_docx(SAMPLE)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "測試社 進度彙整" in text and "第一點" in text and "粗體" in text
    assert doc.tables[0].cell(1, 1).text == "先做預算表"


@pytest.mark.skipif(not pdf_available(), reason="沒有中文字型")
def test_pdf_is_generated():
    data = to_pdf(SAMPLE)
    assert data.startswith(b"%PDF") and len(data) > 1000


def test_schedule_times_follow_minutes():
    rows = [AgendaRow("追蹤", 10), AgendaRow("討論", 45), AgendaRow("臨時動議", 15)]
    spans = [s.span for s in schedule(rows, time(18, 50))]
    assert spans == ["18:50–19:00", "19:00–19:45", "19:45–20:00"]
    md = schedule_markdown(date(2026, 10, 13), time(18, 50), 90, rows, "社辦")
    assert "2026-10-13（二）18:50–20:20" in md and "社辦" in md and "比預定時長少 20 分鐘" in md


@pytest.mark.parametrize("minutes", [30, 60, 90, 120])
def test_template_fills_meeting_length(minutes):
    assert total_minutes(template_rows(minutes, ["行銷", "公關"])) == minutes
