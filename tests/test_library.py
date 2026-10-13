import io
import zipfile

import pytest

from club_agent.agents import ClubQA, club_retriever
from club_agent.library import CATEGORIES, add_file, delete_file, extract_text, file_bytes, file_chunks, list_files
from club_agent.meetings import extract_text as meeting_text
from club_agent.schemas import MeetingAnswer
from club_agent.store import LocalClubStore, StoreError
from club_agent.supabase_store import SupabaseClubStore, migrate_local

from conftest import FakeLLM
from fake_supabase import FakeClient


def make_docx(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_pdf(text: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, text)
    return bytes(pdf.output())


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "budget"
    wb.active.append(["item", "amount"])
    wb.active.append(["poster", 1200])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def make_pptx() -> bytes:
    buffer = io.BytesIO()
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    with zipfile.ZipFile(buffer, "w") as z:
        z.writestr("ppt/slides/slide2.xml", f"<p:sld {ns}><a:t>Second slide</a:t></p:sld>")
        z.writestr("ppt/slides/slide1.xml", f"<p:sld {ns}><a:t>Welcome</a:t><a:t>freshmen</a:t></p:sld>")
    return buffer.getvalue()


def test_extract_text_formats():
    assert "場地要兩週前申請" in extract_text("規定.docx", make_docx("場地要兩週前申請"))
    assert "Venue rules" in extract_text("rules.pdf", make_pdf("Venue rules"))
    assert "poster\t1200" in extract_text("budget.xlsx", make_xlsx())
    assert extract_text("deck.pptx", make_pptx()) == "[第 1 頁] Welcome freshmen\n[第 2 頁] Second slide"
    assert extract_text("note.txt", "社課週四".encode("utf-8")) == "社課週四"
    assert extract_text("photo.png", b"\x89PNG") == "" and extract_text("broken.pdf", b"not a pdf") == ""
    assert "Meeting notes" in meeting_text("minutes.pdf", make_pdf("Meeting notes"))  # 會議記錄也能上傳 PDF
    with pytest.raises(ValueError, match="讀不到文字"):
        meeting_text("scan.pdf", b"not a pdf")


@pytest.mark.parametrize("backend", ["local", "supabase"])
def test_add_list_download_delete(tmp_path, club, backend):
    store = LocalClubStore(tmp_path) if backend == "local" else SupabaseClubStore(FakeClient())
    cid = store.create_club("lib-club", "secret123", club)
    data = make_docx("成果展場地：學活中心")
    f = add_file(store, cid, "成果展 企劃.docx", data, "events", category="企劃書", note="已通過", uploader="小明")
    assert f.key.endswith(".docx") and f.key.isascii() and f.size == len(data) and "學活中心" in f.text
    add_file(store, cid, "海報.png", b"\x89PNG fake", "marketing", category="照片與設計")
    assert len(list_files(store, cid)) == 2
    assert file_bytes(store, cid, f) == data
    assert f.matches("學活") and f.matches("企劃") and not f.matches("贊助")
    delete_file(store, cid, f)
    assert [x.name for x in list_files(store, cid)] == ["海報.png"] and store.get_file(cid, f.key) is None
    with pytest.raises(StoreError, match="空的檔案"):
        add_file(store, cid, "empty.txt", b"", "events")
    with pytest.raises(StoreError, match="20 MB"):
        add_file(store, cid, "big.zip", b"0" * (20 * 1024 * 1024 + 1), "events")
    with pytest.raises(StoreError):
        store.put_file(cid, "../evil.txt", b"x")
    assert "其他" in CATEGORIES


def test_library_feeds_ai_answers(tmp_path, club):
    store = LocalClubStore(tmp_path)
    cid = store.create_club("qa-club", "secret123", club)
    add_file(store, cid, "場地規定.docx", make_docx("借用學活中心演藝廳需在兩週前向課外組申請"), "venue", category="規章與辦法")
    chunks = file_chunks(list_files(store, cid), {"venue": "總務"}.get)
    assert chunks[0].source == "文件庫｜總務｜場地規定.docx"
    llm = FakeLLM({MeetingAnswer: [MeetingAnswer(found=True, answer="兩週前", sources=[])]})
    ClubQA(llm).run("演藝廳要多久前申請？", [], club_retriever([], [], extra=chunks), "（無）")
    assert "兩週前向課外組申請" in llm.calls[0][1]


def test_migrate_includes_files(tmp_path, club):
    local = LocalClubStore(tmp_path)
    cid = local.create_club("mf-club", "secret123", club)
    f = add_file(local, cid, "規章.txt", "社費 500 元".encode("utf-8"), "finance")
    client = FakeClient()
    stats = migrate_local(local, SupabaseClubStore(client), log=lambda _: None)
    assert stats["files"] == 1
    remote = SupabaseClubStore(client)
    assert file_bytes(remote, cid, f) == "社費 500 元".encode("utf-8") and list_files(remote, cid)[0].name == "規章.txt"
