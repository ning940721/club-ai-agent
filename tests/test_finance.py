from datetime import date
from io import BytesIO

import pytest

from club_agent.finance import (
    EXPENSE,
    INCOME,
    PAID,
    FinanceSettings,
    LedgerEntry,
    Reimbursement,
    add_request,
    approve,
    budget_rows,
    get_request,
    get_settings,
    guess_mapping,
    import_rows,
    list_requests,
    mark_paid,
    month_range,
    new_request_code,
    policy_markdown,
    reject,
    report_digest,
    report_markdown,
    review_warnings,
    save_settings,
    summarize_period,
    to_excel,
    transactions,
)
from club_agent.store import LocalClubStore


def make_request(**kw) -> Reimbursement:
    data = dict(applicant="小明", department="events", item="海報輸出", category="文宣印刷", amount=1200,
                invoice_no="AB12345678", receipt_date="2026-10-01", created_at="2026-10-03T10:00:00")
    data.update(kw)
    return Reimbursement(**data)


def test_request_code_format():
    code = new_request_code(date(2026, 10, 6))
    assert code.startswith("R1006-") and len(code) == 9


def test_pin_and_term_defaults():
    fs = FinanceSettings().with_pin("1234")
    assert fs.check_pin("1234") and not fs.check_pin("0000")
    with pytest.raises(ValueError):
        fs.with_pin("12")
    assert FinanceSettings().term(date(2026, 10, 6)) == (date(2026, 8, 1), date(2027, 1, 31))
    assert FinanceSettings().term(date(2027, 1, 10)) == (date(2026, 8, 1), date(2027, 1, 31))
    assert FinanceSettings().term(date(2026, 3, 1)) == (date(2026, 2, 1), date(2026, 7, 31))


def test_review_warnings():
    fs = FinanceSettings(approval_threshold=3000, claim_deadline_days=14)
    ok = make_request()
    assert review_warnings(ok, [ok], fs) == []
    dup = make_request(id="R1003-AAA", invoice_no="ab-1234 5678")
    assert any("可能重複報帳" in w for w in review_warnings(dup, [ok, dup], fs))
    assert not any("可能重複報帳" in w for w in review_warnings(dup, [reject(ok, "財務", "x"), dup], fs))  # 退件的不算
    late = make_request(receipt_date="2026-09-01")
    assert any("超過規定的 14 天" in w for w in review_warnings(late, [late], fs))
    receipt = make_request(invoice_no="收據001")
    assert any("不是統一發票格式" in w for w in review_warnings(receipt, [receipt], fs))
    big = make_request(amount=5000)
    assert any("需要社長同意" in w for w in review_warnings(big, [big], fs))
    assert not any("需要社長同意" in w for w in review_warnings(approve(big, "社長", president_ok=True), [big], fs))


def test_flow_store_and_add_request_never_overwrites(tmp_path, club):
    store = LocalClubStore(tmp_path)
    cid = store.create_club("fin-club", "secret123", club)
    first = add_request(store, cid, make_request(id="R1003-AAA"))
    second = add_request(store, cid, make_request(id="R1003-AAA", item="膠帶"))
    assert first.id != second.id and len(list_requests(store, cid)) == 2
    assert get_request(store, cid, "r1003-aaa ").item == "海報輸出"  # 查詢不分大小寫、忽略空白
    save_settings(store, cid, FinanceSettings(approval_threshold=500))
    assert get_settings(store, cid).approval_threshold == 500


def test_transactions_summary_and_budget():
    paid = mark_paid(approve(make_request(id="P1"), "財務"), date(2026, 10, 5))
    approved = approve(make_request(id="A1", category="餐飲", amount=900, receipt_date="2026-10-02"), "財務")
    pending = make_request(id="Q1", category="餐飲", amount=400)
    rejected = reject(make_request(id="X1", amount=99999), "財務", "重複")
    entries = [
        LedgerEntry(date="2026-09-20", kind=INCOME, category="社費", amount=6000, description="社費"),
        LedgerEntry(date="2026-10-10", kind=EXPENSE, category="雜支", amount=300, description="文具"),
    ]
    txs = transactions([paid, approved, pending, rejected], entries)
    assert [(t.date, t.kind, t.amount, t.ref) for t in txs] == [
        ("2026-09-20", INCOME, 6000, "手動"),
        ("2026-10-05", EXPENSE, 1200, "P1"),
        ("2026-10-10", EXPENSE, 300, "手動"),
    ]
    oct_ = summarize_period(txs, *month_range(2026, 10), opening_balance=1000)
    assert (oct_.opening, oct_.income, oct_.expense, oct_.closing) == (7000, 0, 1500, 5500)
    assert oct_.by_category == {(EXPENSE, "文宣印刷"): 1200, (EXPENSE, "雜支"): 300}

    fs = FinanceSettings(budgets={"餐飲": 1200, "文宣印刷": 1000})
    rows = {b.category: b for b in budget_rows(fs, [paid, approved, pending, rejected], entries, date(2026, 8, 1), date(2027, 1, 31))}
    assert (rows["餐飲"].used, rows["餐飲"].pending, rows["餐飲"].status) == (900, 400, "待審核通過後會超支")
    assert rows["文宣印刷"].status == "超支"
    assert rows["雜支"].status == "未編預算"

    md = report_markdown("測試社", "2026 年 10 月財務月報", oct_, list(rows.values()))
    assert "| 7,000 | 0 | 1,500 | -1,500 | 5,500 |" in md and "| 支出 | 文宣印刷 | 1,200 | 80% |" in md
    digest = report_digest("月報", oct_, list(rows.values()), [paid, approved, pending])
    assert "待審核報帳 1 筆共 400 元；已核准待撥款 1 筆共 900 元" in digest


def test_excel_has_all_sheets():
    from openpyxl import load_workbook

    paid = mark_paid(approve(make_request(), "財務"), date(2026, 10, 5))
    txs = transactions([paid], [LedgerEntry(date="2026-10-01", kind=INCOME, category="社費", amount=3000, description="社費")])
    s = summarize_period(txs, *month_range(2026, 10), opening_balance=0)
    rows = budget_rows(FinanceSettings(budgets={"文宣印刷": 2000}), [paid], [], date(2026, 8, 1), date(2027, 1, 31))
    wb = load_workbook(BytesIO(to_excel("測試社", "10 月月報", s, rows, [paid])))
    assert wb.sheetnames == ["摘要", "收支明細", "類別統計", "預算執行", "報帳明細"]
    assert wb["摘要"]["B8"].value == 1800  # 期末餘額
    assert wb["收支明細"]["F3"].value == 1200 and wb["報帳明細"]["A2"].value == paid.id


def test_policy_uses_settings():
    md = policy_markdown("測試社", FinanceSettings(approval_threshold=2000, claim_deadline_days=10))
    assert "超過 2,000 元" in md and "消費後 10 天內" in md


def test_google_form_import():
    columns = ["時間戳記", "姓名", "購買品項", "金額（元）", "發票號碼", "部門"]
    mapping = guess_mapping(columns)
    assert mapping == {"applicant": "姓名", "item": "購買品項", "amount": "金額（元）", "invoice_no": "發票號碼",
                       "department": "部門", "receipt_date": "時間戳記"}
    rows = [
        {"時間戳記": "2026/9/15 下午 3:20:00", "姓名": "小華", "購買品項": "膠帶", "金額（元）": "$1,250", "發票號碼": "ab 12345678", "部門": "活動組"},
        {"時間戳記": "2026/9/16", "姓名": "", "購買品項": "便當", "金額（元）": "500", "發票號碼": "", "部門": ""},
    ]
    items, skipped = import_rows(rows, mapping, PAID, lambda name: "events" if "活動" in name else "")
    assert len(items) == 1 and skipped == ["第 3 列：缺少申請人、項目或金額"]
    r = items[0]
    assert (r.amount, r.invoice_no, r.receipt_date, r.paid_at, r.department, r.status) == (
        1250, "AB12345678", "2026-09-15", "2026-09-15", "events", PAID)


def test_custom_budget_categories():
    from club_agent.web.finance_pages import budget_from_editor

    categories, budgets = budget_from_editor([
        {"類別": "社課材料", "預算": 3000}, {"類別": " 餐飲 ", "預算": 2000}, {"類別": "", "預算": 500},
        {"類別": "社課材料", "預算": 1000}, {"類別": "保險", "預算": None},
    ])
    assert categories == ["社課材料", "餐飲", "保險"] and budgets == {"社課材料": 4000, "餐飲": 2000}

    fs = FinanceSettings(categories=categories, budgets=budgets)
    old = approve(make_request(category="文宣印刷", amount=800), "財務")  # 已刪掉的類別仍有支出
    rows = budget_rows(fs, [old], [], date(2026, 8, 1), date(2027, 1, 31))
    assert [(b.category, b.budget, b.used, b.status) for b in rows] == [
        ("社課材料", 4000, 0, "正常"), ("餐飲", 2000, 0, "正常"), ("文宣印刷", 0, 800, "未編預算"),
    ]
    assert "社課材料、餐飲、保險" in policy_markdown("測試社", fs)

    columns = ["姓名", "項目", "金額", "類別"]
    rows_in = [{"姓名": "a", "項目": "x", "金額": "100", "類別": "社課材料"}, {"姓名": "b", "項目": "y", "金額": "50", "類別": "其他"}]
    items, _ = import_rows(rows_in, guess_mapping(columns), PAID, lambda _: "", categories=["社課材料", "保險"])
    assert [r.category for r in items] == ["社課材料", "保險"]  # 沒有雜支時歸到最後一個類別
