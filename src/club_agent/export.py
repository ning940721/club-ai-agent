"""把系統產生的 Markdown 報告轉成 Word（.docx）與 PDF，方便幹部直接開啟、列印或轉傳。

只處理報告會用到的 Markdown：標題、段落、項目符號、編號清單、表格、引言與粗體。
PDF 需要中文字型：Windows 使用微軟正黑體；Linux（含 Streamlit Cloud）需安裝 fonts-noto-cjk
（packages.txt 已列出）；也可用環境變數 CLUB_AGENT_PDF_FONT 指定字型檔路徑。
"""

from __future__ import annotations

import io
import logging
import os
import re
from dataclasses import dataclass, field
from pathlib import Path


class ExportError(RuntimeError):
    pass


# ---------------------------------------------------------------------------
# Markdown 解析
# ---------------------------------------------------------------------------


@dataclass
class Block:
    kind: str  # heading / para / bullet / number / quote / table / rule
    text: str = ""
    level: int = 0  # 標題層級，或編號清單的數字
    rows: list[list[str]] = field(default_factory=list)


TABLE_DIVIDER = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")


def _table_cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_markdown(markdown: str) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []
    lines = markdown.replace("\r\n", "\n").split("\n")

    def flush():
        if para:
            blocks.append(Block("para", " ".join(para)))
            para.clear()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            flush()
        elif m := re.match(r"^(#{1,6})\s+(.*)$", stripped):
            flush()
            blocks.append(Block("heading", m.group(2).strip(), level=len(m.group(1))))
        elif stripped.startswith("|") and i + 1 < len(lines) and TABLE_DIVIDER.match(lines[i + 1].strip()):
            flush()
            rows = [_table_cells(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_table_cells(lines[i]))
                i += 1
            width = len(rows[0])
            blocks.append(Block("table", rows=[(r + [""] * width)[:width] for r in rows]))
            continue
        elif m := re.match(r"^[-*+]\s+(.*)$", stripped):
            flush()
            blocks.append(Block("bullet", m.group(1)))
        elif m := re.match(r"^(\d+)[.)]\s+(.*)$", stripped):
            flush()
            blocks.append(Block("number", m.group(2), level=int(m.group(1))))
        elif stripped.startswith(">"):
            flush()
            blocks.append(Block("quote", stripped.lstrip(">").strip()))
        elif re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            flush()
            blocks.append(Block("rule"))
        else:
            para.append(stripped)
        i += 1
    flush()
    return blocks


def inline_runs(text: str) -> list[tuple[str, bool]]:
    """把 **粗體** 拆成 (文字, 是否粗體)；其他 Markdown 記號（*斜體*、`程式碼`）直接去掉。"""
    text = re.sub(r"(?<!\*)\*(?!\*)([^*\n]+?)(?<!\*)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    runs = []
    for n, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            runs.append((part, n % 2 == 1))
    return runs


def plain(text: str) -> str:
    return "".join(t for t, _ in inline_runs(text))


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

DOCX_FONT = "Microsoft JhengHei"


def _set_font(style_or_run, name: str = DOCX_FONT) -> None:
    from docx.oxml.ns import qn

    style_or_run.font.name = name
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


def _add_runs(paragraph, text: str, bold: bool = False) -> None:
    for part, is_bold in inline_runs(text):
        run = paragraph.add_run(part)
        run.bold = bold or is_bold


def to_docx(markdown: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet", "Quote"):
        style = doc.styles[name]
        _set_font(style)
        if name.startswith("Heading") or name == "Title":
            style.font.color.rgb = RGBColor(0x1D, 0x24, 0x33)
    doc.styles["Normal"].font.size = Pt(11)

    first_heading = True
    for b in parse_markdown(markdown):
        if b.kind == "heading":
            if first_heading and b.level == 1:
                doc.add_heading(plain(b.text), level=0)
            else:
                doc.add_heading(plain(b.text), level=min(b.level - 1, 4) or 1)
            first_heading = False
        elif b.kind == "para":
            _add_runs(doc.add_paragraph(), b.text)
        elif b.kind == "bullet":
            _add_runs(doc.add_paragraph(style="List Bullet"), b.text)
        elif b.kind == "number":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)
            _add_runs(p, f"{b.level}.\t{b.text}")
        elif b.kind == "quote":
            _add_runs(doc.add_paragraph(style="Quote"), b.text)
        elif b.kind == "rule":
            doc.add_paragraph("")
        elif b.kind == "table":
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Table Grid"
            for r, row in enumerate(b.rows):
                for c, cell_text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _add_runs(cell.paragraphs[0], cell_text, bold=r == 0)
            doc.add_paragraph("")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

# (一般字型, 粗體字型)；粗體找不到時用一般字型代替
PDF_FONT_CANDIDATES = [
    ("C:/Windows/Fonts/msjh.ttc", "C:/Windows/Fonts/msjhbd.ttc"),
    ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ("/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc", "/usr/share/fonts/noto-cjk/NotoSansCJK-Bold.ttc"),
    ("/System/Library/Fonts/PingFang.ttc", ""),
]
# Noto Sans CJK 字型集合的順序是 JP、KR、SC、TC、HK，取繁體中文（TC）
TTC_INDEX = {"NotoSansCJK": 3}


def find_pdf_font() -> tuple[str, str] | None:
    custom = os.environ.get("CLUB_AGENT_PDF_FONT")
    candidates = [(custom, "")] if custom else PDF_FONT_CANDIDATES
    for regular, bold in candidates:
        if regular and Path(regular).exists():
            return regular, bold if bold and Path(bold).exists() else regular
    return None


def pdf_available() -> bool:
    return find_pdf_font() is not None


def _ttc_index(path: str) -> int:
    return next((i for prefix, i in TTC_INDEX.items() if Path(path).name.startswith(prefix)), 0)


def _display_width(text: str) -> int:
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)


def _column_widths(rows: list[list[str]]) -> tuple[int, ...]:
    """依每欄最長內容分配欄寬：短欄（編號、時間）窄、長欄（理由、目標）寬。"""
    return tuple(max(10, min(60, max(_display_width(plain(r[c])) for r in rows))) for c in range(len(rows[0])))


def to_pdf(markdown: str) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    font = find_pdf_font()
    if font is None:
        raise ExportError("這台主機沒有中文字型，無法產生 PDF，請改下載 Word 檔")
    logging.getLogger("fontTools").setLevel(logging.ERROR)

    pdf = FPDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    regular, bold = font
    pdf.add_font("cjk", "", regular, collection_font_number=_ttc_index(regular))
    pdf.add_font("cjk", "B", bold, collection_font_number=_ttc_index(bold))
    pdf.add_page()
    line_h = 6.2
    sizes = {1: 18, 2: 14, 3: 12.5, 4: 11.5}

    def write_runs(text: str, size: float = 10.5, indent: float = 0, prefix: str = "") -> None:
        left = pdf.l_margin
        pdf.set_x(left + indent)
        if prefix:
            pdf.set_font("cjk", "", size)
            pdf.write(line_h, prefix)
        pdf.set_left_margin(pdf.get_x())  # 換行時對齊文字開頭，而不是清單符號
        for part, is_bold in inline_runs(text):
            pdf.set_font("cjk", "B" if is_bold else "", size)
            pdf.write(line_h, part)
        pdf.set_left_margin(left)
        pdf.ln(line_h + 1.2)

    for b in parse_markdown(markdown):
        if b.kind == "heading":
            size = sizes.get(b.level, 11)
            pdf.ln(2 if b.level > 2 else 4)
            pdf.set_font("cjk", "B", size)
            pdf.multi_cell(0, size * 0.5, plain(b.text), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1.5)
        elif b.kind == "para":
            write_runs(b.text)
        elif b.kind == "bullet":
            write_runs(b.text, indent=3, prefix="•  ")
        elif b.kind == "number":
            write_runs(b.text, indent=3, prefix=f"{b.level}.  ")
        elif b.kind == "quote":
            pdf.set_text_color(90, 98, 112)
            write_runs(b.text, indent=4)
            pdf.set_text_color(0, 0, 0)
        elif b.kind == "rule":
            pdf.ln(3)
        elif b.kind == "table":
            pdf.set_font("cjk", "", 9.5)
            heading_style = FontFace(family="cjk", emphasis="BOLD", fill_color=(241, 239, 234))
            with pdf.table(col_widths=_column_widths(b.rows), line_height=5.4, padding=1.4, headings_style=heading_style, first_row_as_headings=True) as table:
                for row in b.rows:
                    cells = table.row()
                    for cell_text in row:
                        cells.cell(plain(cell_text))
            pdf.ln(3)
    return bytes(pdf.output())
