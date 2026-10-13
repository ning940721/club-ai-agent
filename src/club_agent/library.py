"""文件庫：各部門上傳的檔案（Word、PDF、Excel、簡報、圖片…），依部門與類別整理。

- 檔案內容存在儲存層的檔案空間（本機 data/clubs/<id>/files，線上為 Supabase Storage），
  檔案資訊存在 files 集合。
- 上傳時擷取文字（Word、PDF、Excel、簡報、純文字），用於搜尋與 AI Agent 問答；圖片沒有文字。
- 所有登入社團帳號的幹部都看得到文件庫，財務單據與個資請不要上傳。
"""

from __future__ import annotations

import io
import re
import uuid
import zipfile
from datetime import datetime
from pathlib import PurePath
from xml.etree import ElementTree

from pydantic import BaseModel, Field

from .metrics import decode_csv_bytes
from .retriever import Chunk
from .store import StoreError

COLLECTION = "files"
CATEGORIES = ("會議記錄", "企劃書", "報告", "簡報", "表單與範本", "照片與設計", "規章與辦法", "其他")
MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 60_000
CHUNK_CHARS = 700
CONTENT_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".txt": "text/plain", ".md": "text/markdown", ".csv": "text/csv",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp",
    ".zip": "application/zip",
}


class LibraryFile(BaseModel):
    id: str = Field(default_factory=lambda: uuid.uuid4().hex[:12])
    name: str = Field(description="原始檔名")
    title: str = Field(default="", description="顯示名稱；空白時用檔名")
    department: str
    category: str = "其他"
    note: str = ""
    size: int = 0
    content_type: str = "application/octet-stream"
    uploaded_at: str = Field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))
    uploader: str = ""
    source: str = Field(default="上傳", description="上傳、會議記錄…")
    text: str = Field(default="", description="擷取的文字（搜尋與 AI 問答用）")

    @property
    def key(self) -> str:
        """檔案在儲存空間的代號：id＋副檔名（只用英數字，避免中文路徑問題）。"""
        ext = re.sub(r"[^A-Za-z0-9.]", "", PurePath(self.name).suffix.lower())[:10]
        return f"{self.id}{ext}"

    @property
    def display(self) -> str:
        return self.title or self.name

    def size_label(self) -> str:
        if self.size >= 1024 * 1024:
            return f"{self.size / 1024 / 1024:.1f} MB"
        return f"{max(1, round(self.size / 1024))} KB"

    def matches(self, query: str) -> bool:
        q = query.strip().lower()
        return not q or any(q in x.lower() for x in (self.name, self.title, self.note, self.text))


# ---------------------------------------------------------------------------
# 擷取文字
# ---------------------------------------------------------------------------


def _xml_text(data: bytes, tag: str) -> list[str]:
    root = ElementTree.fromstring(data)
    return [t.text for t in root.iter() if t.tag.endswith(tag) and t.text]


def extract_text(name: str, data: bytes) -> str:
    """擷取檔案中的文字；不支援的格式（例如圖片）或讀取失敗時回傳空字串。"""
    ext = PurePath(name).suffix.lower()
    try:
        if ext in (".txt", ".md", ".csv"):
            text = decode_csv_bytes(data)
        elif ext == ".docx":
            from .meetings import docx_text

            text = docx_text(data)
        elif ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:80])
        elif ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:10]:
                lines.append(f"[{ws.title}]")
                for row in ws.iter_rows(max_row=2000, values_only=True):
                    cells = [str(c) for c in row if c not in (None, "")]
                    if cells:
                        lines.append("\t".join(cells))
            text = "\n".join(lines)
        elif ext == ".pptx":
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                                key=lambda n: int(re.findall(r"\d+", n)[-1]))
                text = "\n".join(f"[第 {i} 頁] " + " ".join(_xml_text(z.read(n), "}t")) for i, n in enumerate(slides, 1))
        else:
            return ""
    except Exception:  # 檔案損毀或加密時，仍然可以存檔，只是不能搜尋內文
        return ""
    return re.sub(r"[ \t]+\n", "\n", text).strip()[:MAX_TEXT]


# ---------------------------------------------------------------------------
# 存取
# ---------------------------------------------------------------------------


def add_file(store, club_id: str, name: str, data: bytes, department: str, *, title: str = "", category: str = "其他",
             note: str = "", uploader: str = "", source: str = "上傳") -> LibraryFile:
    if not data:
        raise StoreError(f"「{name}」是空的檔案")
    if len(data) > MAX_BYTES:
        raise StoreError(f"「{name}」超過 {MAX_BYTES // 1024 // 1024} MB，請壓縮或拆成小檔案")
    entry = LibraryFile(
        name=name, title=title.strip(), department=department, category=category if category in CATEGORIES else "其他",
        note=note.strip(), size=len(data), content_type=CONTENT_TYPES.get(PurePath(name).suffix.lower(), "application/octet-stream"),
        uploader=uploader.strip(), source=source, text=extract_text(name, data),
    )
    store.put_file(club_id, entry.key, data, entry.content_type)
    store.put_doc(club_id, COLLECTION, entry.id, entry.model_dump())
    return entry


def list_files(store, club_id: str) -> list[LibraryFile]:
    """由新到舊。"""
    return sorted((LibraryFile(**d) for d in store.list_docs(club_id, COLLECTION)), key=lambda f: f.uploaded_at, reverse=True)


def update_file(store, club_id: str, entry: LibraryFile) -> None:
    store.put_doc(club_id, COLLECTION, entry.id, entry.model_dump())


def delete_file(store, club_id: str, entry: LibraryFile) -> None:
    store.delete_file(club_id, entry.key)
    store.delete_doc(club_id, COLLECTION, entry.id)


def file_bytes(store, club_id: str, entry: LibraryFile) -> bytes:
    return store.get_file(club_id, entry.key) or b""


def file_chunks(files: list[LibraryFile], dept_name=lambda k: k) -> list[Chunk]:
    """把文件內容切成段落，讓 AI Agent 問答也能從文件庫找答案。"""
    chunks: list[Chunk] = []
    for f in files:
        if not f.text:
            continue
        source = f"文件庫｜{dept_name(f.department)}｜{f.display}"
        buf, size = [], 0
        for line in f.text.splitlines():
            buf.append(line)
            size += len(line)
            if size >= CHUNK_CHARS:
                chunks.append(Chunk(source=source, heading=f.uploaded_at[:10], text="\n".join(buf)))
                buf, size = [], 0
        if buf:
            chunks.append(Chunk(source=source, heading=f.uploaded_at[:10], text="\n".join(buf)))
    return chunks
