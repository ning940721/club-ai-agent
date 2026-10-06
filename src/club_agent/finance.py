"""財務：報帳申請 → 財務審核 → 撥款，收支帳簿、預算執行與月報／學期報表（可下載 Excel）。

資料存在社團的三個集合：
- reimbursements：報帳申請（各部門填寫，財務審核）
- ledger：手動記錄的收入與支出（社費、補助、贊助、非報帳支出…）
- finance：財務設定（財務密碼、核准門檻、預算、學期期間、期初餘額）

已撥款的報帳會自動算進支出，不需要再手動記一次帳。
財務頁面以「財務密碼」上鎖，只讓財務與社長查看；這是同一個社團帳號內的權限區隔，
不是獨立帳號，正式上線（線上資料庫）時再改為個人帳號權限。
"""

from __future__ import annotations

import hashlib
import hmac
import io
import re
import secrets
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timedelta

from pydantic import BaseModel, Field

REIMBURSEMENTS, LEDGER, SETTINGS_COLLECTION, SETTINGS_ID = "reimbursements", "ledger", "finance", "settings"

PENDING, APPROVED, PAID, REJECTED = "待審核", "已核准", "已撥款", "退件"
STATUSES = (PENDING, APPROVED, PAID, REJECTED)
EXPENSE_CATEGORIES = ("活動", "器材設備", "文宣印刷", "餐飲", "交通", "場地", "講師費", "雜支")
INCOME_CATEGORIES = ("社費", "學校補助", "贊助", "活動收入", "其他收入")
PAYMENT_METHODS = ("現金", "轉帳")
INCOME, EXPENSE = "收入", "支出"
MIN_PIN_LENGTH = 4

INVOICE_PATTERN = re.compile(r"^[A-Z]{2}\d{8}$")


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def new_request_code(today: date | None = None) -> str:
    """報帳編號，例：R1006-K7Q（日期＋亂碼，方便口頭查詢）。"""
    today = today or date.today()
    alphabet = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # 去掉容易看錯的 0/O、1/I
    return f"R{today:%m%d}-{''.join(secrets.choice(alphabet) for _ in range(3))}"


def normalize_invoice(text: str) -> str:
    return re.sub(r"[\s\-－]", "", text or "").upper()


class Reimbursement(BaseModel):
    id: str = Field(default_factory=new_request_code)
    created_at: str = Field(default_factory=_now)
    applicant: str
    department: str
    item: str = Field(description="購買項目")
    category: str = "雜支"
    amount: int = Field(gt=0)
    invoice_no: str = Field(default="", description="統一發票號碼；收據可空白並在備註說明")
    receipt_date: str = Field(description="消費日期 YYYY-MM-DD")
    activity: str = Field(default="", description="所屬活動，例：成果展")
    payment_method: str = "現金"
    note: str = ""
    status: str = PENDING
    review_note: str = ""
    reviewed_at: str = ""
    reviewed_by: str = ""
    president_approved: bool = False
    paid_at: str = Field(default="", description="撥款日期 YYYY-MM-DD")
    source: str = "系統申請"

    @property
    def committed(self) -> bool:
        """已核准或已撥款：算進預算使用。"""
        return self.status in (APPROVED, PAID)


class LedgerEntry(BaseModel):
    id: str = Field(default_factory=lambda: uuid.uuid4().hex[:12])
    date: str
    kind: str = Field(description="收入或支出")
    category: str
    amount: int = Field(gt=0)
    description: str
    handler: str = Field(default="", description="經手人")
    invoice_no: str = ""
    created_at: str = Field(default_factory=_now)


class FinanceSettings(BaseModel):
    pin_salt: str = ""
    pin_hash: str = ""
    approval_threshold: int = Field(default=3000, ge=0, description="此金額以上需社長同意")
    claim_deadline_days: int = Field(default=14, ge=1, description="消費後幾天內要申請")
    opening_balance: int = Field(default=0, description="期初餘額（學期開始時帳上的錢）")
    term_start: str = ""
    term_end: str = ""
    budgets: dict[str, int] = Field(default_factory=dict, description="支出類別 → 本學期預算")

    @property
    def has_pin(self) -> bool:
        return bool(self.pin_hash)

    def check_pin(self, pin: str) -> bool:
        return self.has_pin and hmac.compare_digest(_hash_pin(pin, self.pin_salt), self.pin_hash)

    def with_pin(self, pin: str) -> FinanceSettings:
        if len(pin) < MIN_PIN_LENGTH:
            raise ValueError(f"財務密碼至少需要 {MIN_PIN_LENGTH} 個字元")
        salt = secrets.token_hex(16)
        return self.model_copy(update={"pin_salt": salt, "pin_hash": _hash_pin(pin, salt)})

    def term(self, today: date) -> tuple[date, date]:
        """本學期期間；沒設定時用常見的學期（2–7 月、8–隔年 1 月）。"""
        try:
            return date.fromisoformat(self.term_start), date.fromisoformat(self.term_end)
        except ValueError:
            if 2 <= today.month <= 7:
                return date(today.year, 2, 1), date(today.year, 7, 31)
            year = today.year if today.month >= 8 else today.year - 1
            return date(year, 8, 1), date(year + 1, 1, 31)


def _hash_pin(pin: str, salt: str) -> str:
    return hashlib.pbkdf2_hmac("sha256", pin.encode("utf-8"), bytes.fromhex(salt), 200_000).hex()


# ---------------------------------------------------------------------------
# 儲存
# ---------------------------------------------------------------------------


def get_settings(store, club_id: str) -> FinanceSettings:
    data = store.get_doc(club_id, SETTINGS_COLLECTION, SETTINGS_ID)
    return FinanceSettings(**data) if data else FinanceSettings()


def save_settings(store, club_id: str, settings: FinanceSettings) -> None:
    store.put_doc(club_id, SETTINGS_COLLECTION, SETTINGS_ID, settings.model_dump())


def save_request(store, club_id: str, r: Reimbursement) -> None:
    store.put_doc(club_id, REIMBURSEMENTS, r.id, r.model_dump())


def add_request(store, club_id: str, r: Reimbursement) -> Reimbursement:
    """新增申請；編號重複時重新產生，避免覆蓋別人的申請。"""
    while store.get_doc(club_id, REIMBURSEMENTS, r.id):
        r = r.model_copy(update={"id": new_request_code()})
    save_request(store, club_id, r)
    return r


def get_request(store, club_id: str, request_id: str) -> Reimbursement | None:
    data = store.get_doc(club_id, REIMBURSEMENTS, request_id.strip().upper())
    return Reimbursement(**data) if data else None


def list_requests(store, club_id: str) -> list[Reimbursement]:
    """由新到舊。"""
    return sorted((Reimbursement(**d) for d in store.list_docs(club_id, REIMBURSEMENTS)), key=lambda r: r.created_at, reverse=True)


def save_entry(store, club_id: str, e: LedgerEntry) -> None:
    store.put_doc(club_id, LEDGER, e.id, e.model_dump())


def delete_entry(store, club_id: str, entry_id: str) -> None:
    store.delete_doc(club_id, LEDGER, entry_id)


def list_entries(store, club_id: str) -> list[LedgerEntry]:
    return sorted((LedgerEntry(**d) for d in store.list_docs(club_id, LEDGER)), key=lambda e: (e.date, e.created_at))


# ---------------------------------------------------------------------------
# 審核
# ---------------------------------------------------------------------------


def review_warnings(r: Reimbursement, others: list[Reimbursement], settings: FinanceSettings) -> list[str]:
    """審核時要注意的地方：發票格式、重複報帳、逾期申請、需要社長同意。"""
    warnings = []
    invoice = normalize_invoice(r.invoice_no)
    if not invoice:
        warnings.append("沒有填發票號碼；若是收據，請確認備註有說明並收到收據正本")
    elif not INVOICE_PATTERN.match(invoice):
        warnings.append(f"發票號碼「{r.invoice_no}」不是統一發票格式（2 個英文字母＋8 位數字），請確認是否為收據")
    else:
        dup = [o.id for o in others if o.id != r.id and o.status != REJECTED and normalize_invoice(o.invoice_no) == invoice]
        if dup:
            warnings.append(f"同一張發票已經在 {'、'.join(dup)} 報帳過，可能重複報帳")
    try:
        gap = (date.fromisoformat(r.created_at[:10]) - date.fromisoformat(r.receipt_date)).days
        if gap > settings.claim_deadline_days:
            warnings.append(f"消費後 {gap} 天才申請，超過規定的 {settings.claim_deadline_days} 天")
        if gap < 0:
            warnings.append("消費日期晚於申請日期，請確認日期是否填錯")
    except ValueError:
        warnings.append("消費日期格式不正確")
    if needs_president(r, settings) and not r.president_approved:
        warnings.append(f"金額 {r.amount:,} 元達 {settings.approval_threshold:,} 元以上，需要社長同意")
    return warnings


def needs_president(r: Reimbursement, settings: FinanceSettings) -> bool:
    return settings.approval_threshold > 0 and r.amount >= settings.approval_threshold


def approve(r: Reimbursement, reviewer: str, note: str = "", president_ok: bool = False) -> Reimbursement:
    return r.model_copy(
        update={"status": APPROVED, "reviewed_by": reviewer, "reviewed_at": _now(), "review_note": note,
                "president_approved": r.president_approved or president_ok}
    )


def reject(r: Reimbursement, reviewer: str, reason: str) -> Reimbursement:
    return r.model_copy(update={"status": REJECTED, "reviewed_by": reviewer, "reviewed_at": _now(), "review_note": reason})


def mark_paid(r: Reimbursement, paid_on: date) -> Reimbursement:
    return r.model_copy(update={"status": PAID, "paid_at": paid_on.isoformat()})


# ---------------------------------------------------------------------------
# 收支與報表
# ---------------------------------------------------------------------------


@dataclass
class Transaction:
    date: str
    kind: str
    category: str
    amount: int
    description: str
    handler: str
    invoice_no: str
    ref: str  # 報帳編號，或「手動」

    @property
    def signed(self) -> int:
        return self.amount if self.kind == INCOME else -self.amount


def transactions(requests: list[Reimbursement], entries: list[LedgerEntry]) -> list[Transaction]:
    """帳簿：手動記錄的收支＋已撥款的報帳（以撥款日記帳）。"""
    out = [Transaction(e.date, e.kind, e.category, e.amount, e.description, e.handler, e.invoice_no, "手動") for e in entries]
    for r in requests:
        if r.status == PAID:
            desc = f"{r.item}（{r.activity}）" if r.activity else r.item
            out.append(Transaction(r.paid_at or r.receipt_date, EXPENSE, r.category, r.amount, desc, r.applicant, r.invoice_no, r.id))
    return sorted(out, key=lambda t: t.date)


def balance_before(txs: list[Transaction], day: date, opening: int) -> int:
    return opening + sum(t.signed for t in txs if t.date < day.isoformat())


@dataclass
class PeriodSummary:
    start: date
    end: date
    opening: int
    income: int
    expense: int
    by_category: dict[tuple[str, str], int]  # (收入/支出, 類別) → 金額
    items: list[Transaction]

    @property
    def net(self) -> int:
        return self.income - self.expense

    @property
    def closing(self) -> int:
        return self.opening + self.net


def summarize_period(txs: list[Transaction], start: date, end: date, opening_balance: int) -> PeriodSummary:
    items = [t for t in txs if start.isoformat() <= t.date <= end.isoformat()]
    by_category: dict[tuple[str, str], int] = {}
    for t in items:
        by_category[(t.kind, t.category)] = by_category.get((t.kind, t.category), 0) + t.amount
    return PeriodSummary(
        start=start,
        end=end,
        opening=balance_before(txs, start, opening_balance),
        income=sum(t.amount for t in items if t.kind == INCOME),
        expense=sum(t.amount for t in items if t.kind == EXPENSE),
        by_category=dict(sorted(by_category.items())),
        items=items,
    )


def month_range(year: int, month: int) -> tuple[date, date]:
    start = date(year, month, 1)
    nxt = date(year + (month == 12), month % 12 + 1, 1)
    return start, nxt - timedelta(days=1)


@dataclass
class BudgetRow:
    category: str
    budget: int
    used: int  # 已撥款＋已核准待撥款＋手動支出
    pending: int  # 待審核中的金額

    @property
    def remaining(self) -> int:
        return self.budget - self.used

    @property
    def rate(self) -> float:
        return self.used / self.budget if self.budget else 0.0

    @property
    def status(self) -> str:
        if not self.budget:
            return "未編預算" if self.used else "—"
        if self.used > self.budget:
            return "超支"
        if self.used + self.pending > self.budget:
            return "待審核通過後會超支"
        if self.rate >= 0.8:
            return "接近上限"
        return "正常"


def budget_rows(
    settings: FinanceSettings, requests: list[Reimbursement], entries: list[LedgerEntry], start: date, end: date
) -> list[BudgetRow]:
    def in_term(day: str) -> bool:
        return start.isoformat() <= day <= end.isoformat()

    used: dict[str, int] = {}
    pending: dict[str, int] = {}
    for r in requests:
        if not in_term(r.receipt_date):
            continue
        if r.committed:
            used[r.category] = used.get(r.category, 0) + r.amount
        elif r.status == PENDING:
            pending[r.category] = pending.get(r.category, 0) + r.amount
    for e in entries:
        if e.kind == EXPENSE and in_term(e.date):
            used[e.category] = used.get(e.category, 0) + e.amount
    categories = [*EXPENSE_CATEGORIES, *[c for c in [*settings.budgets, *used, *pending] if c not in EXPENSE_CATEGORIES]]
    rows = [BudgetRow(c, settings.budgets.get(c, 0), used.get(c, 0), pending.get(c, 0)) for c in dict.fromkeys(categories)]
    return [r for r in rows if r.budget or r.used or r.pending]


def report_digest(title: str, s: PeriodSummary, budgets: list[BudgetRow], requests: list[Reimbursement]) -> str:
    """給 AI 分析用的文字摘要（不含個人姓名以外的敏感資料）。"""
    lines = [
        f"{title}（{s.start} ~ {s.end}）",
        f"期初餘額 {s.opening:,}；收入 {s.income:,}；支出 {s.expense:,}；期末餘額 {s.closing:,}",
        "類別明細：",
        *[f"- {kind}｜{cat}：{amount:,}" for (kind, cat), amount in s.by_category.items()],
        "預算執行（本學期）：",
        *[f"- {b.category}：預算 {b.budget:,}，已用 {b.used:,}，待審核 {b.pending:,}，{b.status}" for b in budgets],
    ]
    pending = [r for r in requests if r.status == PENDING]
    approved = [r for r in requests if r.status == APPROVED]
    lines.append(f"待審核報帳 {len(pending)} 筆共 {sum(r.amount for r in pending):,} 元；已核准待撥款 {len(approved)} 筆共 {sum(r.amount for r in approved):,} 元")
    big = sorted(s.items, key=lambda t: t.amount, reverse=True)[:5]
    lines += ["金額最大的 5 筆：", *[f"- {t.date}｜{t.kind}｜{t.category}｜{t.description}：{t.amount:,}" for t in big]]
    return "\n".join(lines)


def _cell(text: object) -> str:
    return str(text).replace("|", "｜").replace("\n", " ")


def report_markdown(club_name: str, title: str, s: PeriodSummary, budgets: list[BudgetRow]) -> str:
    out = [
        f"# {club_name} {title}",
        "",
        f"**期間：** {s.start} ~ {s.end}",
        "",
        "## 收支摘要",
        "| 期初餘額 | 收入 | 支出 | 本期結餘 | 期末餘額 |",
        "|---|---|---|---|---|",
        f"| {s.opening:,} | {s.income:,} | {s.expense:,} | {s.net:,} | {s.closing:,} |",
        "",
        "## 類別統計",
        "| 收支 | 類別 | 金額 | 占比 |",
        "|---|---|---|---|",
    ]
    for (kind, cat), amount in s.by_category.items():
        total = s.income if kind == INCOME else s.expense
        out.append(f"| {kind} | {_cell(cat)} | {amount:,} | {amount / total:.0%} |" if total else f"| {kind} | {_cell(cat)} | {amount:,} | — |")
    if budgets:
        out += ["", "## 預算執行（本學期）", "| 類別 | 預算 | 已使用 | 剩餘 | 執行率 | 狀態 |", "|---|---|---|---|---|---|"]
        out += [
            f"| {_cell(b.category)} | {b.budget:,} | {b.used:,} | {b.remaining:,} | {b.rate:.0%} | {b.status} |" for b in budgets
        ]
    out += ["", "## 收支明細", "| 日期 | 收支 | 類別 | 說明 | 金額 | 經手人 | 單據 |", "|---|---|---|---|---|---|---|"]
    out += [
        f"| {t.date} | {t.kind} | {_cell(t.category)} | {_cell(t.description)} | {t.amount:,} | {_cell(t.handler or '—')} | {_cell(t.invoice_no or t.ref)} |"
        for t in s.items
    ]
    if not s.items:
        out.append("| — | — | — | 這段期間沒有收支 | — | — | — |")
    return "\n".join(out) + "\n"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def to_excel(club_name: str, title: str, s: PeriodSummary, budgets: list[BudgetRow], requests: list[Reimbursement]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="F1EFEA")
    thin = Side(style="thin", color="D6D2C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = "#,##0"

    def sheet(ws, headers: list[str], rows: list[list], widths: list[int], money_cols: tuple[int, ...] = ()):
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row > 1 and cell.column in money_cols and isinstance(cell.value, (int, float)):
                    cell.number_format = money
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "摘要"
    ws.append([f"{club_name} {title}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"期間：{s.start} ~ {s.end}"])
    ws.append([])
    for label, value in (("期初餘額", s.opening), ("收入", s.income), ("支出", s.expense), ("本期結餘", s.net), ("期末餘額", s.closing)):
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 2).number_format = money
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 16

    sheet(
        wb.create_sheet("收支明細"),
        ["日期", "收支", "類別", "說明", "收入", "支出", "經手人", "發票／單據", "報帳編號"],
        [[t.date, t.kind, t.category, t.description, t.amount if t.kind == INCOME else None,
          t.amount if t.kind == EXPENSE else None, t.handler, t.invoice_no, "" if t.ref == "手動" else t.ref] for t in s.items],
        [12, 8, 12, 36, 12, 12, 12, 16, 14],
        money_cols=(5, 6),
    )
    sheet(
        wb.create_sheet("類別統計"),
        ["收支", "類別", "金額"],
        [[kind, cat, amount] for (kind, cat), amount in s.by_category.items()],
        [8, 16, 14],
        money_cols=(3,),
    )
    sheet(
        wb.create_sheet("預算執行"),
        ["類別", "預算", "已使用", "待審核", "剩餘", "執行率", "狀態"],
        [[b.category, b.budget, b.used, b.pending, b.remaining, round(b.rate, 4), b.status] for b in budgets],
        [14, 12, 12, 12, 12, 10, 18],
        money_cols=(2, 3, 4, 5),
    )
    for row in wb["預算執行"].iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].number_format = "0%"
    in_period = [r for r in requests if s.start.isoformat() <= r.receipt_date <= s.end.isoformat()]
    sheet(
        wb.create_sheet("報帳明細"),
        ["報帳編號", "申請日", "申請人", "部門", "項目", "類別", "活動", "金額", "發票號碼", "消費日期", "狀態", "審核備註", "撥款日"],
        [[r.id, r.created_at[:10], r.applicant, r.department, r.item, r.category, r.activity, r.amount, r.invoice_no,
          r.receipt_date, r.status, r.review_note, r.paid_at] for r in in_period],
        [12, 11, 10, 10, 26, 10, 14, 10, 14, 11, 9, 20, 11],
        money_cols=(8,),
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 報帳規範建議（依社團設定的門檻與期限產生，可下載給幹部）
# ---------------------------------------------------------------------------


def policy_markdown(club_name: str, settings: FinanceSettings) -> str:
    threshold = f"{settings.approval_threshold:,} 元" if settings.approval_threshold else "（未設定）"
    return f"""# {club_name} 報帳與財務規範

> 由系統依目前的財務設定產生，請依學校核銷規定調整後公告給全體幹部。

## 一、報帳流程
1. **消費前**：超過 {threshold} 的支出，先在幹部群組或幹部會議取得社長同意，再購買。
2. **消費後 {settings.claim_deadline_days} 天內**：到系統「報帳申請」填寫，記下報帳編號。
3. **繳交單據**：發票或收據正本交給財務（可先拍照備份），單據背面寫上報帳編號與申請人。
4. **財務審核**：財務確認金額、單據與預算後核准或退件；退件會註明原因，修正後重新申請。
5. **撥款**：核准後由財務統一撥款，系統標記「已撥款」，申請人可用報帳編號查詢進度。

## 二、單據規定
- 優先索取**統一發票**；發票號碼為 2 個英文字母＋8 位數字（例：AB12345678）。
- 若要向學校核銷，購買時請依學校規定填寫**買受人抬頭與統一編號**，事後無法補開。
- 無法開發票的支出（例：攤商、個人），請開立**收據**，註明日期、品項、金額、收款人簽名，並在報帳備註說明。
- 一張單據只能報帳一次；同一筆消費不可拆成多次申請。

## 三、申請時要填的資料
| 欄位 | 說明 |
|---|---|
| 申請人、部門 | 實際付錢、要收到撥款的人 |
| 項目 | 具體寫出買了什麼，例：「成果展海報輸出 A1 × 5 張」 |
| 類別 | {'、'.join(EXPENSE_CATEGORIES)} |
| 所屬活動 | 例：成果展、期初迎新；日常支出可空白 |
| 金額 | 單據上的實付金額 |
| 發票號碼／消費日期 | 照單據填寫 |
| 撥款方式 | {'、'.join(PAYMENT_METHODS)}（轉帳帳號請私下提供給財務，不要寫在系統裡） |

## 四、財務的例行工作
- **每週**：處理待審核的報帳，核准的款項集中在固定日期撥款。
- **每月**：記錄社費、補助、贊助等收入；核對帳上餘額與實際現金／帳戶是否一致；產出月報給社長。
- **每學期**：期初編列各類別預算並在系統設定；期末產出學期報表，公開收支摘要給社員，交接時移交帳本、單據與報表。

## 五、預算控管
- 各類別預算在系統「財務 → 預算」設定，執行率達 80% 時會提醒。
- 超出預算的支出需先在幹部會議討論，決定從哪個類別調整或動用預備金。
- 建議保留總預算 5%–10% 作為預備金，放在「雜支」類別。
"""


# ---------------------------------------------------------------------------
# 匯入 Google 表單回覆（Google 試算表 → 檔案 → 下載 → CSV）
# ---------------------------------------------------------------------------

FIELD_LABELS = {
    "applicant": "申請人",
    "item": "項目",
    "amount": "金額",
    "invoice_no": "發票號碼",
    "receipt_date": "消費日期",
    "department": "部門",
    "category": "類別",
    "activity": "所屬活動",
    "note": "備註",
}
REQUIRED_FIELDS = ("applicant", "item", "amount")
FIELD_KEYWORDS = {
    "applicant": ("申請人", "姓名", "名字", "填寫人"),
    "item": ("項目", "品項", "品名", "用途", "內容"),
    "amount": ("金額", "費用", "價錢", "總額"),
    "invoice_no": ("發票", "收據", "單據"),
    "receipt_date": ("消費日期", "購買日期", "日期"),
    "department": ("部門", "組別"),
    "category": ("類別", "分類"),
    "activity": ("活動",),
    "note": ("備註", "說明"),
}


def guess_mapping(columns: list[str]) -> dict[str, str]:
    """依欄位名稱猜測對應；Google 表單的「時間戳記」當作消費日期的備用。"""
    mapping: dict[str, str] = {}
    for field, words in FIELD_KEYWORDS.items():
        for col in columns:
            if col not in mapping.values() and any(w in str(col) for w in words):
                mapping[field] = col
                break
    if "receipt_date" not in mapping:
        stamp = next((c for c in columns if "時間戳記" in str(c) or "Timestamp" in str(c)), None)
        if stamp:
            mapping["receipt_date"] = stamp
    return mapping


def parse_amount(value) -> int | None:
    digits = re.sub(r"[^\d.]", "", str(value or ""))
    try:
        amount = round(float(digits))
    except ValueError:
        return None
    return amount if amount > 0 else None


def parse_date(value) -> str | None:
    text = str(value or "").strip()
    match = re.search(r"(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})", text)
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()
    except ValueError:
        return None


def import_rows(
    rows: list[dict], mapping: dict[str, str], status: str, department_lookup, today: date | None = None
) -> tuple[list[Reimbursement], list[str]]:
    """把表單回覆轉成報帳資料；回傳 (成功的資料, 略過的原因)。department_lookup 把部門名稱轉成代號。"""
    today = today or date.today()
    out, skipped = [], []
    for n, row in enumerate(rows, 2):  # 第 1 列是標題
        get = lambda field: str(row.get(mapping.get(field, ""), "") or "").strip()  # noqa: E731
        amount = parse_amount(get("amount"))
        if not get("applicant") or not get("item") or amount is None:
            skipped.append(f"第 {n} 列：缺少申請人、項目或金額")
            continue
        day = parse_date(get("receipt_date")) or today.isoformat()
        category = get("category")
        out.append(
            Reimbursement(
                applicant=get("applicant"), department=department_lookup(get("department")), item=get("item"),
                category=category if category in EXPENSE_CATEGORIES else "雜支", amount=amount,
                invoice_no=normalize_invoice(get("invoice_no")), receipt_date=day, activity=get("activity"),
                note=get("note"), status=status, paid_at=day if status == PAID else "",
                created_at=f"{day}T00:00:00", source="Google 表單匯入",
            )
        )
    return out, skipped
